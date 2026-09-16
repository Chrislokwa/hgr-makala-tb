import io
import json
from datetime import date, datetime, timedelta
from calendar import month_name

from django.contrib import messages
from django.db.models import Count
from django.http import HttpResponse, JsonResponse
from django.shortcuts import render
from django.utils import timezone
from django.views.generic import View

from apps.patients.models import EpisodeTB, IssueFinale, Patient, StatutEpisodeTB, StatutDossier, TypePatient, SiteMaladie, Sexe
from .permissions import StatisticienRequiredMixin
from .services import periode_range, liste_unites, build_dashboard_context


class DashboardView(StatisticienRequiredMixin, View):
    """1. Épidémiologie – Dépistage (page d'accueil statistiques)."""
    template_name = 'statistics/dashboard.html'

    def get(self, request, *args, **kwargs):
        annee = request.GET.get('annee') or str(timezone.localdate().year)
        trimestre = request.GET.get('trimestre') or ''
        mois = request.GET.get('mois') or ''
        unite = request.GET.get('unite') or ''

        if mois:
            trimestre = ''

        debut, fin = periode_range(annee, trimestre if trimestre else None, mois if mois else None)

        context_ctx = build_dashboard_context(debut, fin, unite if unite else None)
        unites = liste_unites()

        context = {
            'active_nav': 'statistics',
            'annee': annee,
            'trimestre': trimestre,
            'mois': mois,
            'unite': unite,
            'debut': debut,
            'fin': fin,
            'unites': unites,
            'ep': context_ctx['epi'],
            'annee_choices': list(range(timezone.localdate().year - 5, timezone.localdate().year + 1)),
        }
        return render(request, self.template_name, context)


class CohorteView(StatisticienRequiredMixin, View):
    """2. Résultats de traitement – Cohorte trimestrielle (page dédiée)."""
    template_name = 'statistics/cohorte.html'

    def get(self, request, *args, **kwargs):
        annee = request.GET.get('annee') or str(timezone.localdate().year)
        trimestre = request.GET.get('trimestre') or ''
        mois = request.GET.get('mois') or ''
        unite = request.GET.get('unite') or ''

        if mois:
            trimestre = ''

        debut, fin = periode_range(annee, trimestre if trimestre else None, mois if mois else None)

        context_ctx = build_dashboard_context(debut, fin, unite if unite else None)
        unites = liste_unites()

        context = {
            'active_nav': 'statistics_cohorte',
            'annee': annee,
            'trimestre': trimestre,
            'mois': mois,
            'unite': unite,
            'debut': debut,
            'fin': fin,
            'unites': unites,
            'cohorte': context_ctx['cohorte'],
            'annee_choices': list(range(timezone.localdate().year - 5, timezone.localdate().year + 1)),
        }
        return render(request, self.template_name, context)


class ExportExcelView(StatisticienRequiredMixin, View):
    """Export Excel pour les rapports de reporting."""

    def get(self, request, *args, **kwargs):
        rapport_id = self.kwargs.get('rapport')
        rapport_info = next((r for r in RAPPORTS if r['id'] == rapport_id), None)
        if not rapport_info:
            return HttpResponse('Rapport introuvable', status=404)

        annee, type_filter, valeur, date_debut, date_fin, debut, fin, periode_label = _get_periode_filters(request)
        qs = _rapport_queryset(rapport_id, debut, fin)

        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        except ImportError:
            return HttpResponse("openpyxl non installé", status=500)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = rapport_id[:31]

        title_font = Font(bold=True, size=14, color="1a6fb0")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="1a6fb0", end_color="1a6fb0", fill_type="solid")
        thin = Side(style="thin", color="dfe3e6")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        center = Alignment(horizontal="center", vertical="center")
        left = Alignment(horizontal="left", vertical="center")

        ws.merge_cells('A1:F1')
        ws['A1'] = "HGR Makala — Programme TB"
        ws['A1'].font = title_font
        ws['A1'].alignment = center

        ws.merge_cells('A2:F2')
        ws['A2'] = rapport_info['titre']
        ws['A2'].font = Font(bold=True, size=12, color="202124")
        ws['A2'].alignment = center

        ws.merge_cells('A3:F3')
        ws['A3'] = f"Période : {periode_label} | Généré le {timezone.localdate().strftime('%d/%m/%Y')}"
        ws['A3'].alignment = center
        ws['A3'].font = Font(italic=True, size=9, color="5f6368")

        headers = ["NDP", "Patient", "Sexe", "Âge", "Date ouverture", "Type"]
        if rapport_id == 'cohorte':
            headers.append("Site")
            headers.append("Statut")
        elif rapport_id == 'en_cours':
            headers.append("Site")
            headers.append("Statut")

        row = 5
        for col, h in enumerate(headers, 1):
            c = ws.cell(row=row, column=col, value=h)
            c.font = header_font
            c.fill = header_fill
            c.alignment = center
            c.border = border

        ws.column_dimensions['A'].width = 16
        ws.column_dimensions['B'].width = 28
        ws.column_dimensions['C'].width = 8
        ws.column_dimensions['D'].width = 8
        ws.column_dimensions['E'].width = 18
        ws.column_dimensions['F'].width = 18
        if len(headers) > 6:
            ws.column_dimensions['G'].width = 14
            ws.column_dimensions['H'].width = 16

        row += 1
        for ep in qs[:500]:
            vals = [
                ep.ndp,
                ep.patient.full_name,
                ep.patient.sexe,
                ep.patient.age,
                ep.date_ouverture.strftime('%d/%m/%Y') if ep.date_ouverture else '',
                ep.get_type_patient_display(),
            ]
            if rapport_id in ('cohorte', 'en_cours'):
                vals.append(ep.get_site_maladie_display())
                vals.append(ep.get_statut_display())
            for col, v in enumerate(vals, 1):
                c = ws.cell(row=row, column=col, value=v)
                c.alignment = center if col in (3, 4) else left
                c.border = border
            row += 1

        if qs.count() > 500:
            ws.cell(row=row, column=1, value=f"... +{qs.count() - 500} autres").font = Font(italic=True, size=9, color="b06000")

        row += 2
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        ws.cell(row=row, column=1, value="Le Responsable du Programme TB").font = Font(bold=True, size=10)
        ws.cell(row=row, column=1).alignment = center
        row += 3
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        ws.cell(row=row, column=1, value="Signature :").font = Font(size=10)
        ws.cell(row=row, column=1).alignment = center

        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        filename = f"{rapport_id}_{debut}_{fin}.xlsx"
        resp = HttpResponse(out.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        resp['Content-Disposition'] = f'attachment; filename="{filename}"'
        return resp


class ExportPdfView(StatisticienRequiredMixin, View):
    """Export PDF pour les rapports de reporting."""

    def get(self, request, *args, **kwargs):
        rapport_id = self.kwargs.get('rapport')
        rapport_info = next((r for r in RAPPORTS if r['id'] == rapport_id), None)
        if not rapport_info:
            return HttpResponse('Rapport introuvable', status=404)

        annee, type_filter, valeur, date_debut, date_fin, debut, fin, periode_label = _get_periode_filters(request)
        qs = _rapport_queryset(rapport_id, debut, fin)

        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.units import mm
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.colors import HexColor
            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
            from reportlab.lib import colors
        except ImportError:
            return HttpResponse("reportlab non installé", status=500)

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4,
                                leftMargin=14*mm, rightMargin=14*mm,
                                topMargin=20*mm, bottomMargin=20*mm,
                                title=rapport_info['titre'], author="HGR Makala")

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle('title', parent=styles['Title'], fontSize=16, leading=20,
                                     textColor=HexColor('#1a6fb0'), alignment=1, spaceAfter=4)
        subtitle_style = ParagraphStyle('subtitle', parent=styles['Normal'], fontSize=11, leading=14,
                                        alignment=1, spaceAfter=2, textColor=HexColor('#202124'))
        info_style = ParagraphStyle('info', parent=styles['Normal'], fontSize=9, leading=12,
                                    alignment=1, textColor=colors.grey, spaceAfter=10)
        footer_style = ParagraphStyle('footer', parent=styles['Normal'], fontSize=9, leading=12,
                                      alignment=1, textColor=colors.grey, spaceBefore=20)

        story = []

        story.append(Paragraph("HGR Makala — Programme Tuberculose", title_style))
        story.append(Paragraph(rapport_info['titre'], subtitle_style))
        story.append(Paragraph(f"Période : {periode_label} &nbsp;|&nbsp; Généré le {timezone.localdate().strftime('%d/%m/%Y')}", info_style))
        story.append(Spacer(1, 6))

        headers = ["NDP", "Patient", "Sexe", "Âge", "Date ouverture", "Type"]
        col_widths = [25*mm, 40*mm, 12*mm, 10*mm, 25*mm, 25*mm]
        if rapport_id in ('cohorte', 'en_cours'):
            headers.extend(["Site", "Statut"])
            col_widths.extend([18*mm, 22*mm])

        data = [headers]
        for ep in qs[:200]:
            row_data = [
                ep.ndp,
                ep.patient.full_name,
                ep.patient.sexe,
                str(ep.patient.age),
                ep.date_ouverture.strftime('%d/%m/%Y') if ep.date_ouverture else '',
                ep.get_type_patient_display(),
            ]
            if rapport_id in ('cohorte', 'en_cours'):
                row_data.append(ep.get_site_maladie_display())
                row_data.append(ep.get_statut_display())
            data.append(row_data)

        if qs.count() > 200:
            data.append(["...", "...", "...", "...", "...", f"+ {qs.count() - 200} autres", "", ""])

        t = Table(data, colWidths=col_widths, repeatRows=1)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), HexColor('#1a6fb0')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ALIGN', (2, 0), (3, -1), 'CENTER'),
            ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#dfe3e6')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f6f8f9')]),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ]))
        story.append(t)

        story.append(Spacer(1, 20))
        story.append(Paragraph(f"Total : {qs.count()} cas", info_style))
        story.append(Spacer(1, 30))

        sig_data = [
            ["", ""],
            ["Le Responsable du Programme TB", "Le Directeur de l'HGR Makala"],
            ["", ""],
            ["_________________________", "_________________________"],
            ["Signature & Cachet", "Signature & Cachet"],
        ]
        sig_table = Table(sig_data, colWidths=[80*mm, 80*mm])
        sig_table.setStyle(TableStyle([
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(sig_table)

        doc.build(story)
        pdf = buffer.getvalue()
        buffer.close()
        filename = f"{rapport_id}_{debut}_{fin}.pdf"
        resp = HttpResponse(pdf, content_type="application/pdf")
        resp['Content-Disposition'] = f'attachment; filename="{filename}"'
        return resp







class TableauBordView(StatisticienRequiredMixin, View):
    """Tableau de bord principal pour l'agent statistique."""
    template_name = 'statistics/tableau_bord.html'

    def get(self, request, *args, **kwargs):
        today = timezone.localdate()

        tous = EpisodeTB.objects.all()

        cas_actifs = tous.filter(
            statut__in=[StatutEpisodeTB.PROVISOIRE, StatutEpisodeTB.CONFIRME, StatutEpisodeTB.EN_TRAITEMENT]
        ).count()

        guerisons = tous.filter(statut=StatutEpisodeTB.CLOTURE, resultat_final='GUERI').count()
        echecs = tous.filter(statut=StatutEpisodeTB.CLOTURE, resultat_final='ECHEC').count()
        abandons = tous.filter(statut=StatutEpisodeTB.CLOTURE, resultat_final='PERDU_DE_VUE').count()
        nouveaux = tous.filter(type_patient=TypePatient.NOUVEAU).count()
        rechutes = tous.filter(type_patient=TypePatient.RECHUTE).count()
        transferes = tous.filter(statut=StatutEpisodeTB.CLOTURE, resultat_final='TRANSFERE').count()
        deces = tous.filter(statut=StatutEpisodeTB.CLOTURE, resultat_final='DECEDE').count()

        total_episodes = tous.count() or 1

        mois_labels = []
        mois_nouveaux = []
        mois_rechutes = []
        mois_abandons = []
        mois_deces = []
        for i in range(11, -1, -1):
            m = today.month - i
            y = today.year
            while m <= 0:
                m += 12
                y -= 1
            debut = date(y, m, 1)
            if m == 12:
                fin = date(y + 1, 1, 1) - timedelta(days=1)
            else:
                fin = date(y, m + 1, 1) - timedelta(days=1)
            base = tous.filter(date_ouverture__gte=debut, date_ouverture__lte=fin)
            mois_labels.append(month_name[m][:3])
            mois_nouveaux.append(base.filter(type_patient=TypePatient.NOUVEAU).count())
            mois_rechutes.append(base.filter(type_patient=TypePatient.RECHUTE).count())
            mois_abandons.append(base.filter(statut=StatutEpisodeTB.CLOTURE, resultat_final='PERDU_DE_VUE').count())
            mois_deces.append(base.filter(statut=StatutEpisodeTB.CLOTURE, resultat_final='DECEDE').count())

        type_labels = []
        type_data = []
        for code, label in TypePatient.choices:
            if code == TypePatient.REPRISE_ABANDON:
                continue
            count = tous.filter(type_patient=code).count()
            type_labels.append(label)
            type_data.append(count)

        sexe_labels = []
        sexe_data = []
        sexe_colors = ['#375a83', '#e74c3c']
        for code, label in Sexe.choices:
            count = Patient.objects.filter(sexe=code).count()
            sexe_labels.append(label)
            sexe_data.append(count)

        derniers = (
            EpisodeTB.objects
            .select_related('patient')
            .order_by('-cree_le')[:5]
        )

        context = {
            'active_nav': 'statistics_tableau_bord',
            'today': today,
            'cas_actifs': cas_actifs,
            'guerisons': guerisons,
            'echecs': echecs,
            'abandons': abandons,
            'nouveaux': nouveaux,
            'rechutes': rechutes,
            'transferes': transferes,
            'deces': deces,
            'total_episodes': total_episodes or 1,
            'mois_labels': json.dumps(mois_labels),
            'mois_nouveaux': json.dumps(mois_nouveaux),
            'mois_rechutes': json.dumps(mois_rechutes),
            'mois_abandons': json.dumps(mois_abandons),
            'mois_deces': json.dumps(mois_deces),
            'type_labels': json.dumps(type_labels),
            'type_data': json.dumps(type_data),
            'sexe_labels': json.dumps(sexe_labels),
            'sexe_data': json.dumps(sexe_data),
            'sexe_colors': json.dumps(sexe_colors),
            'derniers': derniers,
        }
        return render(request, self.template_name, context)


# ──────────────────────────────────────────────────────────────────────────────
# Reporting
# ──────────────────────────────────────────────────────────────────────────────

RAPPORTS = [
    {
        'id': 'nouveaux',
        'titre': 'Liste des nouveaux cas enregistrés',
        'icone': 'file alternate outline',
        'indicateur_label': 'cas enregistrés',
    },
    {
        'id': 'en_cours',
        'titre': 'Liste des cas en cours de traitement',
        'icone': 'clipboard list',
        'indicateur_label': 'cas actifs',
    },
    {
        'id': 'cohorte',
        'titre': 'Rapport des résultats de traitement (cohorte)',
        'icone': 'chart pie',
        'indicateur_label': 'de succès thérapeutique',
        'indicateur_pct': True,
    },
    {
        'id': 'abandons',
        'titre': "Liste des abandons thérapeutiques",
        'icone': 'remove circle outline',
        'indicateur_label': 'abandons',
    },
    {
        'id': 'rechutes',
        'titre': 'Liste des cas de rechute',
        'icone': 'redo',
        'indicateur_label': 'rechutes',
    },
]


def _get_periode_filters(request):
    """Parse les filtres de période depuis la requête.

    Retourne (annee, type_filter, valeur, date_debut, date_fin, debut, fin, periode_label)
    """
    type_filter = request.GET.get('type_filter') or ''
    annee = request.GET.get('annee') or str(timezone.localdate().year)
    valeur = request.GET.get('valeur') or ''
    date_debut_str = request.GET.get('date_debut') or ''
    date_fin_str = request.GET.get('date_fin') or ''

    date_debut = None
    date_fin = None
    if date_debut_str and date_fin_str:
        try:
            date_debut = datetime.strptime(date_debut_str, '%Y-%m-%d').date()
            date_fin = datetime.strptime(date_fin_str, '%Y-%m-%d').date()
        except ValueError:
            pass

    mois = ''
    trimestre = ''
    semestre = ''

    if type_filter == 'mensuel' and valeur:
        mois = valeur
    elif type_filter == 'trimestrielle' and valeur:
        trimestre = valeur
    elif type_filter == 'semestrielle' and valeur:
        semestre = valeur

    debut, fin = periode_range(
        annee,
        trimestre=trimestre or None,
        mois=mois or None,
        semestre=semestre or None,
        date_debut=date_debut,
        date_fin=date_fin,
    )

    periode_label = ''
    if date_debut and date_fin:
        periode_label = f"{date_debut.strftime('%d/%m/%Y')} – {date_fin.strftime('%d/%m/%Y')}"
    elif type_filter == 'semestrielle' and valeur:
        periode_label = f"Semestre {valeur} {annee}"
    elif type_filter == 'trimestrielle' and valeur:
        periode_label = f"Trimestre {valeur} {annee}"
    elif type_filter == 'mensuel' and valeur:
        noms_mois = ['', 'Janvier', 'Février', 'Mars', 'Avril', 'Mai', 'Juin',
                     'Juillet', 'Août', 'Septembre', 'Octobre', 'Novembre', 'Décembre']
        periode_label = f"{noms_mois[int(valeur)]} {annee}"
    else:
        periode_label = f"Année {annee}"

    return annee, type_filter, valeur, date_debut, date_fin, debut, fin, periode_label


def _rapport_queryset(rapport_id, debut, fin, sexe='', site_maladie='', statut='', resultat_final=''):
    qs = EpisodeTB.objects.filter(date_ouverture__gte=debut, date_ouverture__lte=fin)
    if rapport_id == 'nouveaux':
        qs = qs.filter(type_patient=TypePatient.NOUVEAU)
    elif rapport_id == 'en_cours':
        qs = qs.filter(statut__in=[StatutEpisodeTB.PROVISOIRE, StatutEpisodeTB.CONFIRME, StatutEpisodeTB.EN_TRAITEMENT])
        if statut:
            qs = qs.filter(statut=statut)
    elif rapport_id == 'cohorte':
        qs = qs.filter(statut=StatutEpisodeTB.CLOTURE, resultat_final__in=['GUERI', 'TERMINE', 'ECHEC', 'PERDU_DE_VUE', 'DECEDE', 'TRANSFERE'])
        if resultat_final:
            qs = qs.filter(resultat_final=resultat_final)
    elif rapport_id == 'abandons':
        qs = qs.filter(statut=StatutEpisodeTB.CLOTURE, resultat_final='PERDU_DE_VUE')
    elif rapport_id == 'rechutes':
        qs = qs.filter(type_patient=TypePatient.RECHUTE)
    if sexe:
        qs = qs.filter(patient__sexe=sexe)
    if site_maladie:
        qs = qs.filter(site_maladie=site_maladie)
    return qs.select_related('patient')


def _build_rapport_indicateur(rapport_id, qs):
    count = qs.count()
    if rapport_id == 'cohorte':
        total = qs.count() or 1
        succes = qs.filter(resultat_final__in=['GUERI', 'TERMINE']).count()
        return round(succes / total * 100), '%'
    return count, ''


class ReportingView(StatisticienRequiredMixin, View):
    template_name = 'statistics/reporting.html'

    def get(self, request, *args, **kwargs):
        annee, type_filter, valeur, date_debut, date_fin, debut, fin, periode_label = _get_periode_filters(request)

        rapports_data = []
        for r in RAPPORTS:
            qs = _rapport_queryset(r['id'], debut, fin)
            iv, suffixe = _build_rapport_indicateur(r['id'], qs)
            rapports_data.append({
                **r,
                'valeur': iv,
                'suffixe': suffixe,
            })

        context = {
            'active_nav': 'statistics_reporting',
            'annee': annee,
            'type_filter': type_filter,
            'valeur': valeur,
            'date_debut': date_debut.strftime('%Y-%m-%d') if date_debut else '',
            'date_fin': date_fin.strftime('%Y-%m-%d') if date_fin else '',
            'debut': debut,
            'fin': fin,
            'rapports': rapports_data,
            'periode_label': periode_label,
            'annee_choices': list(range(timezone.localdate().year - 5, timezone.localdate().year + 1)),
        }

        if request.headers.get('HX-Request'):
            return render(request, 'statistics/reporting_rapports.html', context)

        return render(request, self.template_name, context)


class RapportDetailView(StatisticienRequiredMixin, View):
    """Page dediee au détail d'un rapport."""
    template_name = 'statistics/rapport_detail.html'
    PER_PAGE = 15

    def get(self, request, *args, **kwargs):
        from django.core.paginator import Paginator, EmptyPage

        rapport_id = kwargs.get('rapport_id')
        rapport_info = next((r for r in RAPPORTS if r['id'] == rapport_id), None)
        if not rapport_info:
            return HttpResponse('Rapport introuvable', status=404)

        annee, type_filter, valeur, date_debut, date_fin, debut, fin, periode_label = _get_periode_filters(request)
        sexe = request.GET.get('sexe', '')
        site_maladie = request.GET.get('site_maladie', '')
        statut = request.GET.get('statut', '')
        resultat_final = request.GET.get('resultat_final', '')
        qs = _rapport_queryset(rapport_id, debut, fin, sexe, site_maladie, statut, resultat_final)
        iv, suffixe = _build_rapport_indicateur(rapport_id, qs)
        total = qs.count()

        paginator = Paginator(qs, self.PER_PAGE)
        page_number = int(request.GET.get('page') or 1)
        try:
            page_obj = paginator.page(page_number)
        except EmptyPage:
            page_obj = paginator.page(paginator.num_pages)

        page_range = []
        cur = page_obj.number
        for num in paginator.page_range:
            if num == cur or (num >= cur - 2 and num <= cur + 2):
                page_range.append(num)

        context = {
            'rapport': rapport_info,
            'items': page_obj,
            'page_obj': page_obj,
            'paginator': paginator,
            'page_range': page_range,
            'valeur': iv,
            'suffixe': suffixe,
            'periode_label': periode_label,
            'total': total,
            'rapport_id': rapport_id,
            'annee': annee,
            'type_filter': type_filter,
            'valeur_filter': valeur,
            'date_debut': date_debut.strftime('%Y-%m-%d') if date_debut else '',
            'date_fin': date_fin.strftime('%Y-%m-%d') if date_fin else '',
            'active_nav': 'statistics_reporting',
            'sexe': sexe,
            'site_maladie': site_maladie,
            'statut': statut,
            'resultat_final': resultat_final,
            'sexe_choices': Sexe.choices,
            'site_maladie_choices': SiteMaladie.choices,
            'statut_en_cours_choices': [
                (StatutEpisodeTB.PROVISOIRE, 'Provisoire'),
                (StatutEpisodeTB.CONFIRME, 'Confirmé'),
                (StatutEpisodeTB.EN_TRAITEMENT, 'En traitement'),
            ],
            'resultat_final_choices': [
                ('GUERI', 'Guéri'),
                ('TERMINE', 'Terminé'),
                ('ECHEC', 'Échec'),
                ('PERDU_DE_VUE', 'Perdu de vue'),
                ('DECEDE', 'Décédé'),
                ('TRANSFERE', 'Transféré'),
            ],
        }

        if request.headers.get('HX-Request'):
            return render(request, 'statistics/rapport_detail_body.html', context)

        return render(request, self.template_name, context)
